"""Export a LegalDocument to .txt, .docx and .pdf with optional branding.

Each exporter takes (doc, branding, logo) where `logo` is the tuple returned by
security.process_logo (png_bytes, width_px, height_px) or None.
"""
import io
import os
import textwrap
from xml.sax.saxutils import escape

from .config import DISCLAIMER
from .schemas import Branding, LegalDocument

DOCX_FONTS = {"sans": "Arial", "serif": "Times New Roman", "mono": "Courier New"}


def _hex_to_rgb(h: str):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def _tint(h: str, amount: float = 0.88) -> str:
    """Lighten a hex colour toward white (for table label backgrounds)."""
    r, g, b = _hex_to_rgb(h)
    return "%02X%02X%02X" % tuple(int(c + (255 - c) * amount) for c in (r, g, b))


def _lines(body: str):
    """Yield (is_bullet, text) for each non-empty line of a section body."""
    for ln in body.replace("\r", "").split("\n"):
        ln = ln.strip()
        if not ln:
            continue
        if ln[:2] in ("- ", "* ", "• "):
            yield True, ln[2:].strip()
        else:
            yield False, ln


def _logo_size_in(logo, max_h=0.65, max_w=2.4):
    _, w, h = logo
    height = max_h
    width = height * w / h
    if width > max_w:
        width, height = max_w, max_w * h / w
    return width, height


# =========================================================================== TXT
def to_txt(doc: LegalDocument, branding: Branding, logo=None) -> bytes:
    W = 80
    out = []
    if branding.company_name:
        out.append(branding.company_name.upper().center(W))
    out += [doc.title.upper().center(W), "=" * W, ""]
    if doc.preamble:
        out += textwrap.wrap(doc.preamble, W) + [""]
    if doc.terms:
        pad = max(len(t.label) for t in doc.terms) + 2
        out.append("KEY TERMS")
        out.append("-" * W)
        for t in doc.terms:
            wrapped = textwrap.wrap(t.value, W - pad) or [""]
            out.append(t.label.ljust(pad) + wrapped[0])
            out += [" " * pad + w for w in wrapped[1:]]
        out += ["-" * W, ""]
    for i, s in enumerate(doc.sections, 1):
        out.append(f"{i}. {s.heading.upper()}")
        for bullet, text in _lines(s.body):
            if bullet:
                out += textwrap.wrap(text, W, initial_indent="   - ", subsequent_indent="     ")
            else:
                out += textwrap.wrap(text, W, initial_indent="   ", subsequent_indent="   ")
            out.append("")
    if doc.signatories:
        out += ["SIGNATURES", "-" * W, ""]
        for s in doc.signatories:
            out += ["_" * 36, s.name, s.role, "Date: " + "_" * 24, ""]
    out += ["-" * W] + textwrap.wrap(DISCLAIMER, W)
    if branding.footer_text:
        out.append(branding.footer_text)
    return ("\n".join(out) + "\n").encode("utf-8")


# =========================================================================== DOCX
def to_docx(doc: LegalDocument, branding: Branding, logo=None) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    font = DOCX_FONTS[branding.font]
    accent = RGBColor(*_hex_to_rgb(branding.accent_color))
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin, sec.bottom_margin = Inches(0.9), Inches(0.9)

    def set_font(style, size, bold=None, color=None):
        style.font.name = font
        style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color is not None:
            style.font.color.rgb = color
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]

    set_font(d.styles["Normal"], 11)
    set_font(d.styles["Title"], 22, True, accent)
    set_font(d.styles["Heading 2"], 12.5, True, accent)
    set_font(d.styles["List Bullet"], 11)
    pbdr = d.styles["Title"].element.pPr.find(qn("w:pBdr")) if d.styles["Title"].element.pPr is not None else None
    if pbdr is not None:                       # drop the default blue rule under the title
        pbdr.getparent().remove(pbdr)
    d.styles["Normal"].paragraph_format.space_after = Pt(6)
    d.styles["Normal"].paragraph_format.line_spacing = 1.12

    def shade(cell, hex6):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex6)
        tcPr.append(shd)

    def bottom_border(par, color_hex):
        pPr = par._p.get_or_add_pPr()
        pb = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color_hex)):
            b.set(qn(k), v)
        pb.append(b); pPr.append(pb)

    def field(par, instr):
        for kind, text in (("begin", None), (None, instr), ("end", None)):
            r = par.add_run()
            if kind:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)

    # ---- header: logo + company name
    hp = sec.header.paragraphs[0]
    if logo:
        w_in, _ = _logo_size_in(logo)
        hp.add_run().add_picture(io.BytesIO(logo[0]), width=Inches(w_in))
    if branding.company_name:
        r = hp.add_run(("   " if logo else "") + branding.company_name)
        r.bold = True; r.font.color.rgb = accent; r.font.size = Pt(12)
    if logo or branding.company_name:
        bottom_border(hp, branding.accent_color.lstrip("#"))

    # ---- footer: custom text + page numbers
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if branding.footer_text:
        fr = fp.add_run(branding.footer_text + "   |   "); fr.font.size = Pt(9)
    pr = fp.add_run("Page "); pr.font.size = Pt(9)
    field(fp, "PAGE")
    for r in fp.runs:
        r.font.size = Pt(9)

    # ---- title / preamble
    t = d.add_paragraph(doc.title, style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if doc.preamble:
        d.add_paragraph(doc.preamble)

    # ---- key terms table
    if doc.terms:
        h = d.add_paragraph(); r = h.add_run("Key Terms"); r.bold = True; r.font.color.rgb = accent
        tbl = d.add_table(rows=0, cols=2)
        tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width, tbl.columns[1].width = Inches(2.0), Inches(4.5)
        for term in doc.terms:
            cells = tbl.add_row().cells
            cells[0].width, cells[1].width = Inches(2.0), Inches(4.5)
            shade(cells[0], _tint(branding.accent_color))
            cells[0].paragraphs[0].add_run(term.label).bold = True
            cells[1].paragraphs[0].add_run(term.value)
            for c in cells:
                c.paragraphs[0].paragraph_format.space_after = Pt(2)
        d.add_paragraph()

    # ---- clauses
    for i, s in enumerate(doc.sections, 1):
        hp2 = d.add_paragraph(f"{i}. {s.heading}", style="Heading 2")
        hp2.paragraph_format.keep_with_next = True
        for bullet, text in _lines(s.body):
            p = d.add_paragraph(text, style="List Bullet" if bullet else "Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if not bullet else WD_ALIGN_PARAGRAPH.LEFT

    # ---- signatures
    if doc.signatories:
        sp = d.add_paragraph(); sp.paragraph_format.space_before = Pt(18); sp.paragraph_format.keep_with_next = True
        r = sp.add_run("Signatures"); r.bold = True; r.font.color.rgb = accent; r.font.size = Pt(12.5)
        d.add_paragraph("The parties have signed this Agreement on the dates written below.").paragraph_format.keep_with_next = True
        tbl = d.add_table(rows=1, cols=len(doc.signatories))
        for cell, s in zip(tbl.rows[0].cells, doc.signatories):
            cell.paragraphs[0].paragraph_format.space_before = Pt(26)
            cell.paragraphs[0].add_run("_" * 30)
            for text, bold in ((s.name, True), (s.role, False), ("Date: " + "_" * 18, False)):
                p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(1)
                p.add_run(text).bold = bold
        tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in tbl.rows[0].cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True

    dp = d.add_paragraph(); dp.paragraph_format.space_before = Pt(18)
    r = dp.add_run(DISCLAIMER); r.italic = True; r.font.size = Pt(8.5)
    d.core_properties.title = doc.title
    d.core_properties.author = branding.company_name or "LegalEase"

    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


# =========================================================================== PDF
_FONT_FILES = {   # family -> (regular, bold, italic) candidate TTF names
    "sans": ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf"),
    "serif": ("DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf"),
    "mono": ("DejaVuSansMono.ttf", "DejaVuSansMono-Bold.ttf", "DejaVuSansMono-Oblique.ttf"),
}
_FONT_DIRS = [os.path.join(os.path.dirname(__file__), "fonts"), "/usr/share/fonts/truetype/dejavu",
              "/usr/share/fonts/dejavu", "/Library/Fonts", "/System/Library/Fonts/Supplemental",
              "C:\\Windows\\Fonts"]
_STD = {"sans": ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique"),
        "serif": ("Times-Roman", "Times-Bold", "Times-Italic"),
        "mono": ("Courier", "Courier-Bold", "Courier-Oblique")}
_registered = {}


def _pdf_fonts(family: str):
    """Prefer bundled/system DejaVu (full Unicode); fall back to PDF built-ins."""
    if family in _registered:
        return _registered[family]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    names, unicode_ok = _STD[family], False
    for folder in _FONT_DIRS:
        paths = [os.path.join(folder, f) for f in _FONT_FILES[family]]
        if all(os.path.exists(p) for p in paths):
            try:
                reg = [f"LE-{family}-{k}" for k in ("r", "b", "i")]
                for n, p in zip(reg, paths):
                    pdfmetrics.registerFont(TTFont(n, p))
                names, unicode_ok = tuple(reg), True
                break
            except Exception:
                continue
    _registered[family] = (names, unicode_ok)
    return _registered[family]


def _pdf_text(text: str, unicode_ok: bool) -> str:
    text = escape(text)
    if unicode_ok:
        return text
    return text.encode("cp1252", "replace").decode("cp1252").replace("\ufffd", "?")


def to_pdf(doc: LegalDocument, branding: Branding, logo=None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    (reg, bold, ital), uni = _pdf_fonts(branding.font)
    T = lambda s: _pdf_text(s, uni)
    accent = colors.HexColor(branding.accent_color)
    light = colors.HexColor("#" + _tint(branding.accent_color))

    base = ParagraphStyle("base", fontName=reg, fontSize=10.5, leading=14.5, alignment=TA_JUSTIFY, spaceAfter=6)
    title = ParagraphStyle("title", parent=base, fontName=bold, fontSize=21, leading=26, alignment=TA_CENTER,
                           textColor=accent, spaceAfter=14)
    h2 = ParagraphStyle("h2", parent=base, fontName=bold, fontSize=12, leading=15, textColor=accent,
                        spaceBefore=10, spaceAfter=4, alignment=0, keepWithNext=1)
    cell = ParagraphStyle("cell", parent=base, alignment=0, spaceAfter=0, fontSize=10, leading=13)
    cellb = ParagraphStyle("cellb", parent=cell, fontName=bold)
    small = ParagraphStyle("small", parent=base, fontName=ital, fontSize=8, leading=10, textColor=colors.grey)

    story = []
    if logo or branding.company_name:
        row = []
        if logo:
            w_in, h_in = _logo_size_in(logo)
            row.append(Image(io.BytesIO(logo[0]), width=w_in * inch, height=h_in * inch))
        if branding.company_name:
            row.append(Paragraph(T(branding.company_name), ParagraphStyle(
                "co", parent=base, fontName=bold, fontSize=13, textColor=accent, alignment=0, spaceAfter=0)))
        if len(row) == 2:
            first = w_in * inch + 14
            widths = [first, 6.5 * inch - first]
        else:
            widths = [6.5 * inch]
        hdr = Table([row], colWidths=widths, hAlign="LEFT")
        hdr.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                 ("LINEBELOW", (0, 0), (-1, -1), 1.2, accent), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story += [hdr, Spacer(1, 16)]

    story.append(Paragraph(T(doc.title), title))
    if doc.preamble:
        story.append(Paragraph(T(doc.preamble), base))

    if doc.terms:
        story += [Spacer(1, 6), Paragraph(T("Key Terms"), ParagraphStyle("kt", parent=h2, spaceBefore=0))]
        rows = [[Paragraph(T(t.label), cellb), Paragraph(T(t.value), cell)] for t in doc.terms]
        tbl = Table(rows, colWidths=[2.0 * inch, 4.5 * inch])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), light), ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8B8B8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story += [tbl, Spacer(1, 8)]

    for i, s in enumerate(doc.sections, 1):
        story.append(Paragraph(T(f"{i}. {s.heading}"), h2))
        bullets = []

        def flush():
            if bullets:
                story.append(ListFlowable([ListItem(Paragraph(T(b), base), leftIndent=14) for b in bullets],
                                          bulletType="bullet", start="\u2022" if uni else "-", leftIndent=14))
                bullets.clear()
        for is_b, text in _lines(s.body):
            if is_b:
                bullets.append(text)
            else:
                flush()
                story.append(Paragraph(T(text), base))
        flush()

    if doc.signatories:
        block = [Spacer(1, 14), Paragraph(T("Signatures"), h2),
                 Paragraph(T("The parties have signed this Agreement on the dates written below."), base)]
        cols = []
        for s in doc.signatories:
            cols.append([Spacer(1, 26), Paragraph("_" * 30, cell), Paragraph(T(s.name), cellb),
                         Paragraph(T(s.role), cell), Paragraph("Date: " + "_" * 18, cell)])
        w = 6.5 * inch / len(cols)
        sig = Table([cols], colWidths=[w] * len(cols))
        sig.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        story.append(KeepTogether(block + [sig, Spacer(1, 16), Paragraph(T(DISCLAIMER), small)]))
    else:
        story += [Spacer(1, 16), Paragraph(T(DISCLAIMER), small)]

    def decorate(canvas, d):
        canvas.saveState()
        canvas.setFont(reg, 8.5)
        canvas.setFillColor(colors.grey)
        footer = (branding.footer_text + "   |   " if branding.footer_text else "") + f"Page {d.page}"
        canvas.drawCentredString(letter[0] / 2, 0.5 * inch, footer.encode("cp1252", "replace").decode("cp1252") if not uni else footer)
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=0.9 * inch,
                      bottomMargin=0.9 * inch, title=doc.title,
                      author=branding.company_name or "LegalEase").build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


EXPORTERS = {
    "txt": (to_txt, "text/plain; charset=utf-8"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": (to_pdf, "application/pdf"),
}
